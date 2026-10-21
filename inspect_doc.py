import docx

doc = docx.Document(r"C:\Users\ASUS\Desktop\uni\TG2079\TG_2079.docx")
print("=== PARAGRAPHS WITH IMAGES ===")
for i, p in enumerate(doc.paragraphs):
    blips = p._p.xpath('.//a:blip')
    if blips:
        print(f"P{i:03d}: {len(blips)} blip(s) | text: '{p.text.strip()}'")

print("\n=== TOTAL TABLES ===", len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols")
