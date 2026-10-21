import docx

doc = docx.Document(r"C:\Users\ASUS\Desktop\uni\TG2079\TG_2079.docx")

print("Document styles present:")
for s in doc.styles:
    if s.type == docx.enum.style.WD_STYLE_TYPE.PARAGRAPH and ('Heading' in s.name or 'Normal' in s.name or 'List' in s.name):
        print(f"  {s.name} (id: {s.style_id})")

table = doc.tables[0]
print("\nTable 0 properties:")
print(f"Style: {table.style.name if table.style else 'None'}")
print(f"Rows: {len(table.rows)}, Cols: {len(table.columns)}")
for r_idx, row in enumerate(table.rows):
    cells = [c.text.strip() for c in row.cells]
    print(f"Row {r_idx}: {cells}")

# Check XML of a Heading 1
for p in doc.paragraphs:
    if p.style.name == 'Heading 1' and p.text.strip():
        print("\nHeading 1 XML snippet:")
        print(p._p.xml[:400])
        break

# Check XML of a Heading 2
for p in doc.paragraphs:
    if p.style.name == 'Heading 2' and p.text.strip():
        print("\nHeading 2 XML snippet:")
        print(p._p.xml[:400])
        break

# Check XML of a List Paragraph
for p in doc.paragraphs:
    if p.style.name == 'List Paragraph' and p.text.strip():
        print("\nList Paragraph XML snippet:")
        print(p._p.xml[:400])
        break
